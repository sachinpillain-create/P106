"""
build_full_guide.py
Assembles and saves the complete PROTOTYPE_MASTER_EXPLANATION_GUIDE.docx document.
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from guide_sections_part1 import build_part1
from guide_sections_part2 import build_part2

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_border(cell, **kwargs):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in('w:tcBorders')
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    for edge in ('top', 'left', 'bottom', 'right'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = f'w:{edge}'
            element = tcBorders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tcBorders.append(element)
            for key in ['val', 'color', 'sz', 'space']:
                if key in edge_data:
                    element.set(qn(f'w:{key}'), str(edge_data[key]))

def add_styled_heading(doc, text, level):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.name = 'Calibri'
    
    if level == 1:
        p.paragraph_format.space_before = Pt(22)
        p.paragraph_format.space_after = Pt(6)
        run.font.size = Pt(16.5)
        run.font.color.rgb = RGBColor(10, 37, 64) # Deep Navy #0A2540
        pPr = p._p.get_or_add_pPr()
        pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="4" w:color="007791"/></w:pBdr>')
        pPr.append(pBdr)
    elif level == 2:
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(30, 58, 138) # Slate Blue #1E3A8A
    elif level == 3:
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(13, 148, 136) # Teal Accent #0D9488
    return p

def add_body_paragraph(doc, text, bold_prefix=None, space_after=4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.name = 'Calibri'
        r_pre.font.size = Pt(10)
        r_pre.font.color.rgb = RGBColor(10, 37, 64)
    r_body = p.add_run(text)
    r_body.font.name = 'Calibri'
    r_body.font.size = Pt(10)
    r_body.font.color.rgb = RGBColor(45, 55, 72)
    return p

def add_bullet_point(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.name = 'Calibri'
        r_pre.font.size = Pt(10)
        r_pre.font.color.rgb = RGBColor(10, 37, 64)
    r_body = p.add_run(text)
    r_body.font.name = 'Calibri'
    r_body.font.size = Pt(10)
    r_body.font.color.rgb = RGBColor(45, 55, 72)
    return p

def add_callout(doc, text_prefix, text_body, title=None, hex_color="0A2540", bg_color="F8FAFC"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(7.0)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=120, bottom=120, left=180, right=160)
    set_cell_border(cell,
                    left={'val': 'single', 'sz': '24', 'color': hex_color},
                    top={'val': 'none'}, right={'val': 'none'}, bottom={'val': 'none'})
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    if title:
        run_t = p.add_run(f"■ {title}\n")
        run_t.bold = True
        run_t.font.name = 'Calibri'
        run_t.font.size = Pt(10.5)
        run_t.font.color.rgb = RGBColor(10, 37, 64)
    if text_prefix:
        run_pre = p.add_run(text_prefix + " ")
        run_pre.bold = True
        run_pre.font.name = 'Calibri'
        run_pre.font.size = Pt(10)
        run_pre.font.color.rgb = RGBColor(10, 37, 64)
    run_b = p.add_run(text_body)
    run_b.font.name = 'Calibri'
    run_b.font.size = Pt(9.5)
    run_b.font.color.rgb = RGBColor(45, 55, 72)
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(4)

def format_table(doc, col_widths, headers, data, align_cols=None):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    hdr_row = tbl.rows[0]
    for i, h_text in enumerate(headers):
        cell = hdr_row.cells[i]
        cell.width = col_widths[i]
        set_cell_background(cell, "0A2540")
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        set_cell_border(cell,
                        bottom={'val': 'single', 'sz': '12', 'color': '007791'},
                        top={'val': 'none'}, left={'val': 'none'}, right={'val': 'none'})
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(h_text)
        run.bold = True
        run.font.name = 'Calibri'
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(255, 255, 255)
        if align_cols and i < len(align_cols):
            p.alignment = align_cols[i]

    for r_idx, row_data in enumerate(data):
        row = tbl.add_row()
        bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, cell_value in enumerate(row_data):
            cell = row.cells[c_idx]
            cell.width = col_widths[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            set_cell_border(cell,
                            bottom={'val': 'single', 'sz': '4', 'color': 'E2E8F0'},
                            top={'val': 'none'}, left={'val': 'none'}, right={'val': 'none'})
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(str(cell_value))
            run.font.name = 'Calibri'
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(45, 55, 72)
            if align_cols and c_idx < len(align_cols):
                p.alignment = align_cols[c_idx]
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(4)
    return tbl

def generate_master_document(output_filename="PROTOTYPE_MASTER_EXPLANATION_GUIDE.docx"):
    print(f"Creating master document: {output_filename}...")
    doc = docx.Document()
    
    # 0.75 inch margins for clean professional typography
    for s in doc.sections:
        s.top_margin = Inches(0.75)
        s.bottom_margin = Inches(0.75)
        s.left_margin = Inches(0.75)
        s.right_margin = Inches(0.75)
        
        # Header setup
        header = s.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("JAALBHEDI (SIH26106) — Prototype Technical & Forensic Master Guide")
        hrun.font.name = 'Calibri'
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(139, 147, 161)
        
        # Footer setup
        footer = s.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        frun = fp.add_run("Smart India Hackathon 2026 — Master Prototype Specification (Confidential)")
        frun.font.name = 'Calibri'
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(139, 147, 161)

    print("Assembling Part 1 (Metadata, Executive Summary, Architecture, 12-Stage Pipeline)...")
    build_part1(doc, add_styled_heading, add_body_paragraph, add_bullet_point, add_callout, format_table)

    print("Assembling Part 2 (Mathematical Models, Codebase, HUD Playbook, Legal, Viva Q&A, Glossary)...")
    build_part2(doc, add_styled_heading, add_body_paragraph, add_bullet_point, add_callout, format_table)

    target_path = os.path.join(os.path.dirname(os.path.dirname(__file__)), output_filename)
    doc.save(target_path)
    print(f"Master guide successfully generated and saved to: {target_path}")
    return target_path

if __name__ == "__main__":
    generate_master_document()
